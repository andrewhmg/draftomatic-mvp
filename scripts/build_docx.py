# Build: IS 581 Week 4 - Draftomatic MVP submission docx (functional-MVP + fallback doc in one)
# Run: /usr/bin/python3 build_mvp_submission.py
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

WD = '/Users/andrewhogge/Documents/School/IS581/mvp-build'
OUT = '/Users/andrewhogge/Downloads/Draftomatic MVP (Week 4).docx'

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)


def h(text, lvl=1):
    doc.add_heading(text, level=lvl)


def p(text, bold_lead=None):
    par = doc.add_paragraph()
    if bold_lead:
        r = par.add_run(bold_lead + ' ')
        r.bold = True
    par.add_run(text)
    return par


def bullet(text, bold_lead=None):
    par = doc.add_paragraph(style='List Bullet')
    if bold_lead:
        r = par.add_run(bold_lead + ' ')
        r.bold = True
    par.add_run(text)


def shot(name, caption):
    path = os.path.join(WD, 'screens', name)
    doc.add_picture(path, width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------- Title ----------
t = doc.add_paragraph()
r = t.add_run('Draftomatic — MVP Development')
r.bold = True
r.font.size = Pt(20)
sub = doc.add_paragraph()
r = sub.add_run('IS 581 · Week 4 · Andrew Hogge')
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
link = doc.add_paragraph()
r = link.add_run('Live working demo (use it directly): ')
r.bold = True
r2 = link.add_run('https://draftomatic-mvp.vercel.app')
r2.font.color.rgb = RGBColor(0x1a, 0x56, 0xdb)

h('What Draftomatic is', 1)
p('Draftomatic is the autopilot blog for small-business websites, built exactly to the scope '
  'defined in my Week 3 MVP Description. The owner connects their site once, approves an '
  'AI-generated content plan, and from then on drafts are written, covered with a social '
  'image, and published to their CMS on a weekly schedule. The owner\u2019s total time cost is '
  'about five minutes a week; that number is the value proposition.')
p('The critical path is Connect \u2192 Plan \u2192 Generate \u2192 Publish, plus the trust controls from '
  'the Week 3 story map (draft preview, autopilot toggle with pause, publish status log). '
  'The single core assumption the MVP tests: small-business owners will let AI publish to '
  'their site with minimal review \u2014 in product terms, they leave autopilot on.')

h('What is built and working', 1)
bullet('sign-up fields, site connection, and automatic industry detection from the site\u2019s '
       'own content (a live page is scanned when reachable; a seed company backs the demo).', 'Connect:')
bullet('twelve SEO-ranked article topics generated from the detected industry, each with a '
       'target keyword; approve-all in one tap, or select, edit, delete, and add topics individually.', 'Plan:')
bullet('a complete draft per topic \u2014 headline, slug, meta title, meta description, and a '
       '600\u20131,000-word body in seven sections \u2014 and one-click regenerate produces a genuinely '
       'different draft.', 'Generate:')
bullet('each draft builds a Framer CMS payload with a documented field mapping plus an '
       'auto-generated 1200\u00d7630 cover image; every publish attempt is logged with its status '
       'and attempt number, and failed publishes retry from the dashboard.', 'Publish:')
bullet('a preview screen shows the full draft before anything goes live; autopilot publishes '
       'automatically on the weekly schedule and can be paused with one toggle. All state '
       'persists across server restarts.', 'Trust:')

h('The four screens', 1)
p('The MVP is exactly the four screens specified in Week 3 \u2014 nothing more.')
shot('01-connect.png', 'Screen 1 — Connect: sign-up, connect the Framer site, industry auto-detected.')
shot('02-topics.png', 'Screen 2 — Topics: ranked content plan with target keywords; one-tap approve.')
shot('03-preview.png', 'Screen 3 — Preview: full draft with SEO meta, 1200\u00d7630 cover, approve \u2192 publish.')
shot('04-dashboard.png', 'Screen 4 — Dashboard: autopilot toggle, published posts, status log with retry.')

h('How it was verified', 1)
p('The app runs with a single command (python3 server.py) using only the Python standard '
  'library \u2014 no installs, no API keys. Twenty automated checks passed covering the full flow: '
  'health, all four pages, connect with industry detection, topic generation and approval, '
  'draft generation with the 600-word minimum enforced, regenerate producing a different '
  'draft, publish writing the CMS payload with cover and status-log entry, and the autopilot '
  'toggle persisting across a real server restart. Each screen was then captured live in a '
  'browser at 1440\u00d7900 (the four figures above are those captures).')

h('What is real vs. demo', 1)
bullet('the connected site is a seed company (a landscaping business), not my own site.')
bullet('article text comes from a template engine, not a live LLM, so the demo runs with no '
       'API keys; the generator interface is the same one an LLM backend would plug into.')
bullet('publishing writes the exact Framer CMS payload it would send, instead of calling '
       'Framer\u2019s API. The live path is implemented and activates automatically when a Framer '
       'API key is present; it has not been exercised against a real Framer account yet.')

h('Remaining steps and timeline', 1)
p('Week 1 \u2014 go live on one real site:')
bullet('create a Framer CMS collection and validate the field mapping end-to-end with a real '
       'API key (the publish call itself is already built and wired).')
bullet('swap the template generator for an LLM backend behind the same generate/regenerate '
       'interface.')
bullet('add the weekly digest email ("3 articles published this week").')
p('Week 2 \u2014 put it in front of owners:')
bullet('run 3\u20135 small-business owners through the live autopilot for one publishing cycle.')
bullet('measure the core assumption directly: how many leave autopilot on by default. Target: '
       'at least 3 of 5 still on after two weekly cycles.')

h('Bottom line', 1)
p('A functional MVP is ready for testing today: the full Connect \u2192 Plan \u2192 Generate \u2192 '
  'Publish loop runs end-to-end, with the minimum features from the Week 3 definition and '
  'nothing else. It is deliberately imperfect \u2014 per the rubric, minimality is the point \u2014 '
  'and it is sufficient for a real owner to use and evaluate this week.')

doc.save(OUT)
print('saved', OUT, os.path.getsize(OUT), 'bytes')